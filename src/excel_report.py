"""Executive Excel report stage: KPI summary + data quality scorecard workbook.

Produces a single .xlsx with two sheets aimed at a non-technical executive
audience: a KPI Summary and a Data Quality Summary with red/yellow/green
conditional formatting so data-quality issues are visible at a glance.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.metrics_engine import MetricResult
from src.validate import DataQualityReport

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

UNIT_FORMATS = {
    "percent": '0.00"%"',
    "currency": '"$"#,##0.00',
    "count": "#,##0",
}


def _style_header_row(ws: Worksheet, row: int, num_cols: int) -> None:
    """Apply the standard header styling to a row."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autofit_columns(ws: Worksheet, widths: list[int]) -> None:
    """Set column widths to reasonable fixed values (openpyxl has no true autofit)."""
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_kpi_sheet(ws: Worksheet, metrics: list[MetricResult]) -> None:
    """Populate the KPI Summary sheet."""
    ws.title = "KPI Summary"
    ws["A1"] = "Carmari Academy — Advancement KPI Summary"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")

    headers = ["Metric", "Value", "Unit", "Version", "Formula", "Notes"]
    header_row = 3
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)
    _style_header_row(ws, header_row, len(headers))

    for r, metric in enumerate(metrics, start=header_row + 1):
        ws.cell(row=r, column=1, value=metric.name)
        value_cell = ws.cell(row=r, column=2, value=metric.value)
        value_cell.number_format = UNIT_FORMATS.get(metric.unit, "General")
        ws.cell(row=r, column=3, value=metric.unit)
        ws.cell(row=r, column=4, value=metric.version)
        ws.cell(row=r, column=5, value=metric.formula).alignment = Alignment(wrap_text=True)
        ws.cell(row=r, column=6, value=metric.notes).alignment = Alignment(wrap_text=True)

    _autofit_columns(ws, [28, 14, 10, 10, 55, 50])
    ws.freeze_panes = f"A{header_row + 1}"


def _status_fill(severity: str, status: str) -> PatternFill:
    """Pick a red/yellow/green fill for a validation rule's severity + status."""
    if status == "pass":
        return GREEN_FILL
    if severity == "critical":
        return RED_FILL
    return YELLOW_FILL


def _write_dq_sheet(ws: Worksheet, dq_report: DataQualityReport) -> None:
    """Populate the Data Quality Summary sheet."""
    ws.title = "Data Quality Summary"
    ws["A1"] = "Carmari Academy — Data Quality Scorecard"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:G1")
    ws["A2"] = (
        f"Overall status: {dq_report.overall_status.upper()}  |  "
        f"{dq_report.critical_failures} critical issue(s), {dq_report.warning_failures} warning(s)  |  "
        f"{dq_report.total_records:,} total records"
    )
    ws["A2"].font = Font(italic=True)

    headers = ["Rule ID", "Description", "Severity", "Status", "Records Affected", "% Affected", "Sample IDs"]
    header_row = 4
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)
    _style_header_row(ws, header_row, len(headers))

    for r, rule in enumerate(dq_report.rule_results, start=header_row + 1):
        ws.cell(row=r, column=1, value=rule.rule_id)
        ws.cell(row=r, column=2, value=rule.description).alignment = Alignment(wrap_text=True)
        ws.cell(row=r, column=3, value=rule.severity.upper())
        status_cell = ws.cell(row=r, column=4, value=rule.status.upper())
        ws.cell(row=r, column=5, value=rule.records_affected)
        ws.cell(row=r, column=6, value=rule.pct_affected / 100).number_format = "0.00%"
        ws.cell(row=r, column=7, value=", ".join(rule.sample_ids[:5]))

        fill = _status_fill(rule.severity, rule.status)
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).fill = fill
        status_cell.font = Font(bold=True)

    _autofit_columns(ws, [22, 55, 12, 10, 16, 12, 40])
    ws.freeze_panes = f"A{header_row + 1}"


def generate_excel_report(
    metrics: list[MetricResult], dq_report: DataQualityReport, output_path: str | Path
) -> Path:
    """Build the two-sheet executive Excel workbook.

    Args:
        metrics: Computed KPI results.
        dq_report: Data quality validation results.
        output_path: Destination path for the .xlsx file.

    Returns:
        The path the workbook was written to.
    """
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    kpi_sheet = wb.active
    _write_kpi_sheet(kpi_sheet, metrics)

    dq_sheet = wb.create_sheet()
    _write_dq_sheet(dq_sheet, dq_report)

    wb.save(path)
    return path
